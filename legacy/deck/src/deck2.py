# -*- coding: utf-8 -*-
"""Light editorial deck engine, built on the user's Slides_template.pptx."""
import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

TPL = "legacy/deck/assets/Slides_template.pptx"
def C(h): return RGBColor(int(h[0:2],16), int(h[2:4],16), int(h[4:6],16))
BG      = C("F0F4F8")
INK     = C("16324F")
BODY    = C("33556E")
MUTED   = C("8CA0B3")
BLUE    = C("2E6E9E")
BLUE_D  = C("143D63")
BLUE_M  = C("6FA3C7")
BLUE_L  = C("BBD3E5")
ORANGE  = C("E8763A")
TINT_B  = C("E3EDF5")
TINT_O  = C("FBEBE0")
RULE    = C("C3D2DE")
WHITE   = C("FFFFFF")
NAVY    = C("0E2841")
FONT    = "Calibri"
SW, SH  = 13.3333, 7.5
LEFT, RIGHT = 0.68, 12.66
CW = RIGHT - LEFT

# ------------------------------------------------------------------ basics
def clone(prs, idx):
    src = prs.slides[idx]
    dst = prs.slides.add_slide(src.slide_layout)
    for sh in list(dst.shapes):
        sh._element.getparent().remove(sh._element)
    for sh in src.shapes:
        dst.shapes._spTree.append(copy.deepcopy(sh._element))
    return dst

def to_back(shape):
    sp = shape._element
    tree = sp.getparent()
    tree.remove(sp)
    tree.insert(2, sp)

def backdrop(slide, color=BG):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(SW), Inches(SH))
    s.fill.solid(); s.fill.fore_color.rgb = color
    s.line.fill.background(); s.shadow.inherit = False
    to_back(s)
    return s

def box(slide, l, t, w, h):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tb.text_frame.word_wrap = True
    return tb

def para(tf, text, size=14, bold=False, color=BODY, align=PP_ALIGN.LEFT,
         first=False, italic=False, before=0, after=5, indent=0):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    r = p.add_run(); r.text = text
    f = r.font; f.name = FONT; f.size = Pt(size); f.bold = bold; f.italic = italic
    f.color.rgb = color
    p.alignment = align
    p.space_before = Pt(before); p.space_after = Pt(after)
    if indent: p.level = indent
    return p

def rich(tf, chunks, size=14, align=PP_ALIGN.LEFT, first=False, before=0, after=5):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    for txt, bold, col in chunks:
        r = p.add_run(); r.text = txt
        f = r.font; f.name = FONT; f.size = Pt(size); f.bold = bold; f.color.rgb = col
    p.alignment = align; p.space_before = Pt(before); p.space_after = Pt(after)
    return p

def rule(slide, l, t, w, color=RULE, h=0.011):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = color
    s.line.fill.background(); s.shadow.inherit = False
    return s

# ------------------------------------------------------------------ chrome
def chrome(slide, tagtxt, author, title=None, lead=None):
    ref = list(slide.shapes)          # capture BEFORE the backdrop shifts indices
    backdrop(slide)
    sh = ref[3]; sh.width = Inches(3.6)
    tf = sh.text_frame; tf.clear()
    para(tf, tagtxt, size=13, color=WHITE, first=True, after=0)
    tf2 = ref[1].text_frame; tf2.clear()
    para(tf2, author, size=13, color=WHITE, first=True, after=0)
    t = ref[4]
    t.left, t.top, t.width, t.height = Inches(LEFT), Inches(0.26), Inches(CW), Inches(0.70)
    tt = t.text_frame; tt.clear()
    if title is not None:
        sz = 27 if len(title) <= 62 else 23
        para(tt, title, size=sz, bold=True, color=INK, first=True, after=0)
    if lead:
        lb = box(slide, LEFT, 1.00, CW, 0.42)
        para(lb.text_frame, lead, size=14.5, color=BODY, first=True, after=0)
    return slide

def content_slide(prs, tagtxt, author, title, lead=None):
    s = clone(prs, 2)
    return chrome(s, tagtxt, author, title, lead)

def section_slide(prs, tagtxt, author, num, title, sub):
    s = clone(prs, 2)
    ref = list(s.shapes)
    backdrop(s)
    sh = ref[3]; sh.width = Inches(3.6)
    tf = sh.text_frame; tf.clear(); para(tf, tagtxt, size=13, color=WHITE, first=True, after=0)
    tf2 = ref[1].text_frame; tf2.clear()
    para(tf2, author, size=13, color=WHITE, first=True, after=0)
    t = ref[4]; t.text_frame.clear()
    t.left, t.top, t.width, t.height = Inches(LEFT), Inches(0.26), Inches(CW), Inches(0.4)
    nb = box(s, LEFT, 2.68, 4.0, 0.36)
    para(nb.text_frame, num, size=15, bold=True, color=BLUE_M, first=True, after=0)
    tb = box(s, LEFT, 3.06, 11.4, 0.78)
    para(tb.text_frame, title, size=34, bold=True, color=INK, first=True, after=0)
    rule(s, LEFT, 3.96, 6.6, BLUE_M, 0.030)
    sb = box(s, LEFT, 4.14, 10.8, 0.5)
    para(sb.text_frame, sub, size=15, color=MUTED, first=True, after=0)
    return s

# ------------------------------------------------------------------ pieces
def picture(slide, path, top=1.52, bottom=6.90, max_w=12.0):
    iw, ih = Image.open(path).size
    ah, aw = bottom - top, max_w
    w = aw; h = w * ih / iw
    if h > ah: h = ah; w = h * iw / ih
    return slide.shapes.add_picture(path, Inches((SW - w) / 2),
                                    Inches(top + (ah - h) / 2), Inches(w), Inches(h))

def rule_table(slide, data, l, t, w, col_w=None, size=13, head_size=13,
               row_h=0.42, head_h=0.40, aligns=None, accent=None):
    """Example-style table: horizontal rules only, no cell fills."""
    rows, cols = len(data), len(data[0])
    h = head_h + row_h * (rows - 1)
    gf = slide.shapes.add_table(rows, cols, Inches(l), Inches(t), Inches(w), Inches(h))
    tbl = gf.table
    tbl.first_row = False; tbl.horz_banding = False
    # replace the inherited "Medium Style 2" with "No Style, No Grid"
    from pptx.oxml.ns import qn
    tblPr = tbl._tbl.find(qn("a:tblPr"))
    if tblPr is not None:
        for el in tblPr.findall(qn("a:tableStyleId")):
            tblPr.remove(el)
        sid = tblPr.makeelement(qn("a:tableStyleId"), {})
        sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
        tblPr.append(sid)
    if col_w:
        tot = sum(col_w)
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = Emu(int(Inches(w) * cw / tot))
    tbl.rows[0].height = Inches(head_h)
    for r in range(1, rows): tbl.rows[r].height = Inches(row_h)
    for r in range(rows):
        for c in range(cols):
            cell = tbl.cell(r, c)
            cell.fill.background()
            cell.margin_left = Inches(0.02 if c == 0 else 0.10)
            cell.margin_right = Inches(0.10)
            cell.margin_top = Inches(0.04); cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            txt = data[r][c]; col = None
            if isinstance(txt, tuple): txt, col = txt
            tf = cell.text_frame; tf.clear(); tf.word_wrap = True
            p = tf.paragraphs[0]
            run = p.add_run(); run.text = str(txt)
            f = run.font; f.name = FONT
            f.size = Pt(head_size if r == 0 else size)
            f.bold = (r == 0)
            f.color.rgb = INK if r == 0 else (col or BODY)
            if aligns and c < len(aligns):
                p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER,
                               "r": PP_ALIGN.RIGHT}[aligns[c]]
    rule(slide, l, t - 0.02, w, INK, 0.019)
    rule(slide, l, t + head_h - 0.01, w, INK, 0.010)
    rule(slide, l, t + h + 0.01, w, INK, 0.010)
    return gf

def callout(slide, l, t, w, label, body, color=BLUE_D, tint=TINT_B, bh=0.36,
            body_h=0.62, size=13.5, label_size=13.5):
    b = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(bh))
    b.fill.solid(); b.fill.fore_color.rgb = color
    b.line.fill.background(); b.shadow.inherit = False
    tf = b.text_frame; tf.word_wrap = True
    tf.margin_left = Inches(0.14); tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para(tf, label, size=label_size, bold=True, color=WHITE, first=True, after=0)
    d = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t + bh),
                               Inches(w), Inches(body_h))
    d.fill.solid(); d.fill.fore_color.rgb = tint
    d.line.fill.background(); d.shadow.inherit = False
    tf2 = d.text_frame; tf2.word_wrap = True
    tf2.margin_left = Inches(0.14); tf2.margin_right = Inches(0.14)
    tf2.margin_top = Inches(0.08); tf2.vertical_anchor = MSO_ANCHOR.TOP
    for i, ln in enumerate(body if isinstance(body, list) else [body]):
        para(tf2, ln, size=size, color=BODY, first=(i == 0), after=3)
    return d

def kpi_row(slide, items, t=1.55, h=1.05, l=LEFT, w=CW, gap=0.20):
    n = len(items); cw = (w - gap * (n - 1)) / n
    for i, (v, lab) in enumerate(items):
        x = l + i * (cw + gap)
        rule(slide, x, t, cw, BLUE_M, 0.030)
        tb = box(slide, x, t + 0.10, cw, h - 0.10)
        tf = tb.text_frame
        para(tf, v, size=30, bold=True, color=BLUE_D, first=True, after=2)
        para(tf, lab, size=12.5, color=MUTED, after=0)

def bullets(slide, l, t, w, h, head, items, size=14, head_size=15, gap=6, color=INK):
    tb = box(slide, l, t, w, h)
    tf = tb.text_frame
    para(tf, head, size=head_size, bold=True, color=color, first=True, after=8)
    for it in items:
        p = tf.add_paragraph()
        r0 = p.add_run(); r0.text = "•   "
        r0.font.name = FONT; r0.font.size = Pt(size); r0.font.color.rgb = BLUE_M
        r = p.add_run(); r.text = it
        r.font.name = FONT; r.font.size = Pt(size); r.font.color.rgb = BODY
        p.space_after = Pt(gap)
    return tb

def note(slide, text, t=6.55, size=12, color=MUTED, l=LEFT, w=CW):
    tb = box(slide, l, t, w, 0.40)
    para(tb.text_frame, text, size=size, color=color, first=True, after=0)
    return tb
