# -*- coding: utf-8 -*-
"""Short Stage-1 deck: what the final df is, and what built it."""
import sys; sys.path.insert(0, "legacy/deck/src")
from deck2 import *
from content4 import C
from pptx.enum.text import PP_ALIGN


def reorder_prune(prs, order):
    lst = prs.slides._sldIdLst
    ids = list(lst)
    keep = [ids[i] for i in order]
    for i in ids:
        if i not in keep:
            prs.part.drop_rel(i.rId)
        lst.remove(i)
    for i in keep:
        lst.append(i)


def build(lang):
    T = C[lang]
    F = lambda n: f"legacy/deck/fig4/{n}_{lang}.png"
    prs = Presentation(TPL)
    AU = prs.slides[2].shapes[1].text_frame.text
    N = 0

    # ---------------------------------------------------------- title
    s = prs.slides[0]
    tf = s.shapes[1].text_frame; tf.clear()
    p = tf.paragraphs[0]; r = p.add_run(); r.text = T["title"]
    r.font.name = FONT; r.font.size = Pt(46); r.font.bold = True; r.font.color.rgb = NAVY
    p.alignment = PP_ALIGN.CENTER
    p2 = tf.add_paragraph(); r2 = p2.add_run(); r2.text = T["sub"]
    r2.font.name = FONT; r2.font.size = Pt(16); r2.font.color.rgb = BLUE
    p2.alignment = PP_ALIGN.CENTER; p2.space_before = Pt(12)
    tk = s.shapes[4].text_frame; tk.clear()
    para(tk, T["kicker"], size=18, bold=True, color=BODY, align=PP_ALIGN.CENTER,
         first=True, after=0)
    s.shapes[2].top, s.shapes[2].height = Inches(4.16), Inches(0.96)
    gb = s.shapes[2].text_frame; gb.clear()
    for i, line in enumerate(T["group"]):
        para(gb, line, size=17, bold=True, color=INK, align=PP_ALIGN.CENTER,
             first=(i == 0), after=3)
    nb = box(s, 1.6, 5.24, 10.1, 0.36)
    para(nb.text_frame, T["group_note"], size=13, color=MUTED,
         align=PP_ALIGN.CENTER, first=True, after=0)
    tg = s.shapes[5].text_frame; tg.clear()
    para(tg, T["tag"], size=12, color=WHITE, align=PP_ALIGN.CENTER, first=True, after=0)

    def new(title, lead=None):
        nonlocal N
        N += 1
        return content_slide(prs, T["tag"], AU, title, lead)

    # ---------------------------------------------------- 1. one row
    s = new(T["s1_t"])
    qb = box(s, LEFT, 1.20, CW, 0.70)
    para(qb.text_frame, T["s1_q"], size=26, bold=True, color=BLUE_D, first=True, after=0)
    y = 2.24
    for i, (k, v) in enumerate(T["s1_key"]):
        x = LEFT + i * 4.16
        rule(s, x, y, 3.86, BLUE_M, 0.030)
        tb = box(s, x, y + 0.14, 3.86, 0.80)
        para(tb.text_frame, k, size=19, bold=True, color=INK, first=True, after=3)
        para(tb.text_frame, v, size=13.5, color=MUTED, after=0)
    rule(s, LEFT, 3.60, CW, RULE, 0.008)
    kpi_row(s, T["s1_kpi"], t=3.90, h=1.10)
    note(s, T["s1_n"], t=5.40, size=13.5)

    # ---------------------------------------------------- 2. blocks
    s = new(T["s2_t"])
    picture(s, F("blocks"), 1.24, 6.32, 11.6)
    note(s, T["s2_n"], t=6.44, size=13.5)

    # ---------------------------------------------------- 3. sources
    s = new(T["s3_t"])
    picture(s, F("sources"), 1.24, 6.32, 12.4)
    note(s, T["s3_n"], t=6.44, size=13.5)

    # ---------------------------------------------------- 4. source -> columns
    s = new(T["s4_t"])
    rule_table(s, T["s4_tbl"], LEFT, 1.24, CW, col_w=[4.6, 6.0, 1.3],
               size=14, head_size=14, row_h=0.50, head_h=0.44,
               aligns=["l", "l", "r"])

    # ---------------------------------------------------- 5. survival label
    s = new(T["s5_t"])
    rule_table(s, T["s5_tbl"], LEFT, 1.40, CW, col_w=[3.4, 8.6],
               size=15, head_size=15, row_h=0.72, head_h=0.46)
    callout(s, LEFT, 5.42, CW, T["s5_c"][0], T["s5_c"][1], BLUE_D, TINT_B,
            bh=0.36, body_h=0.62, size=14)

    # ---------------------------------------------------- 6. target label
    s = new(T["s6_t"], T["s6_l"])
    picture(s, F("labels"), 1.54, 4.44, 11.4)
    rule(s, LEFT, 4.62, CW, RULE, 0.008)
    kpi_row(s, T["s6_kpi"], t=4.86, h=1.04)
    callout(s, LEFT, 6.10, 7.20, T["s6_c"][0], T["s6_c"][1], BLUE_D, TINT_B,
            bh=0.34, body_h=0.48, size=13)
    note(s, T["s6_n"], t=6.30, l=8.16, w=4.50, size=13)

    # ---------------------------------------------------------- thanks
    ty = prs.slides[3]
    ref = list(ty.shapes)
    backdrop(ty)
    sh = ref[3]; sh.width = Inches(3.6)
    tfx = sh.text_frame; tfx.clear()
    para(tfx, T["tag"], size=13, color=WHITE, first=True, after=0)
    tfa = ref[1].text_frame; tfa.clear()
    para(tfa, AU, size=13, color=WHITE, first=True, after=0)
    t4 = ref[4]
    t4.left, t4.top, t4.width, t4.height = Inches(0), Inches(2.85), Inches(SW), Inches(0.9)
    tt = t4.text_frame; tt.clear()
    para(tt, T["thanks"], size=42, bold=True, color=INK, align=PP_ALIGN.CENTER,
         first=True, after=0)
    rule(ty, 5.17, 4.02, 3.0, BLUE_M, 0.030)
    sb = box(ty, 1.6, 4.24, 10.1, 0.6)
    para(sb.text_frame, T["thanks_sub"], size=14, color=MUTED,
         align=PP_ALIGN.CENTER, first=True, after=0)

    reorder_prune(prs, [0] + list(range(4, 4 + N)) + [3])
    out = f"legacy/deck/out/Final_DF_{'VI' if lang == 'vi' else 'EN'}.pptx"
    prs.save(out)
    print(f"saved {out}  ({N + 2} slides)")


if __name__ == "__main__":
    for lg in ("vi", "en"):
        build(lg)
