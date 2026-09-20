# -*- coding: utf-8 -*-
"""Flag shapes that physically collide on a slide (tables measured at true height)."""
import sys
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import ImageFont, ImageDraw, Image
import glob

CAR = (glob.glob("/usr/share/fonts/**/Carlito-Regular.ttf", recursive=True) or
       ["/usr/share/fonts/truetype/freefont/FreeSans.ttf"])[0]
CARB = (glob.glob("/usr/share/fonts/**/Carlito-Bold.ttf", recursive=True) or
        ["/usr/share/fonts/truetype/freefont/FreeSansBold.ttf"])[0]
_d = ImageDraw.Draw(Image.new("RGB", (10, 10)))
_fc = {}
def fnt(sz, b):
    k = (round(sz), b)
    if k not in _fc: _fc[k] = ImageFont.truetype(CARB if b else CAR, max(6, round(sz)))
    return _fc[k]

def I(v): return Emu(v).inches if v is not None else 0.0

def wrapped_lines(text, size, bold, width_in):
    f = fnt(size * 8, bold)               # 8 px per pt for measuring
    maxw = width_in * 72 * 8
    n = 0
    for raw in text.split("\n"):
        words = raw.split(" "); line = ""
        for w in words:
            t = (line + " " + w).strip()
            if _d.textlength(t, font=f) <= maxw or not line: line = t
            else: n += 1; line = w
        n += 1
    return max(1, n)

def table_true_height(shape):
    tbl = shape.table
    cw = [I(c.width) for c in tbl.columns]
    tot = 0.0
    for ri, row in enumerate(tbl.rows):
        need = I(row.height)
        for ci, cell in enumerate(row.cells):
            txt = cell.text
            if not txt.strip(): continue
            sz, bold = 12.0, False
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.size: sz = r.font.size.pt
                    bold = bold or bool(r.font.bold)
            ml = I(cell.margin_left) + I(cell.margin_right)
            lines = wrapped_lines(txt, sz, bold, max(.4, cw[ci] - ml))
            h = lines * sz * 1.22 / 72 + I(cell.margin_top) + I(cell.margin_bottom)
            need = max(need, h)
        tot += need
    return tot

def rects(slide):
    out = []
    for sh in slide.shapes:
        x, y, w, h = I(sh.left), I(sh.top), I(sh.width), I(sh.height)
        if w > 13 and h > 7: continue                      # backdrop
        if y > 7.0: continue                               # footer chrome
        kind = "shape"
        if sh.has_table:
            h = table_true_height(sh); kind = "table"
        elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE: kind = "pic"
        elif sh.has_text_frame and sh.text_frame.text.strip(): kind = "text"
        elif h < .06: continue                             # rules
        else: kind = "block"
        out.append((x, y, w, h, kind, sh.name[:22],
                    (sh.text_frame.text[:34] if sh.has_text_frame else "")))
    return out

def check(path):
    prs = Presentation(path); bad = 0
    for i, sl in enumerate(prs.slides):
        rs = rects(sl)
        for a in range(len(rs)):
            for b in range(a + 1, len(rs)):
                x1, y1, w1, h1, k1, n1, t1 = rs[a]
                x2, y2, w2, h2, k2, n2, t2 = rs[b]
                ox = max(0, min(x1 + w1, x2 + w2) - max(x1, x2))
                oy = max(0, min(y1 + h1, y2 + h2) - max(y1, y2))
                if ox <= .02 or oy <= .04: continue
                if {k1, k2} == {"text"} and ox * oy < .12: continue
                small = min(w1 * h1, w2 * h2)
                if ox * oy / max(small, .01) < .12: continue
                print(f"  slide {i+1}: {k1}[{t1 or n1}] x {k2}[{t2 or n2}] "
                      f"overlap {ox:.2f}x{oy:.2f}in")
                bad += 1
        for x, y, w, h, k, n, t in rs:
            if y + h > 7.05:
                print(f"  slide {i+1}: {k}[{t or n}] runs to {y+h:.2f}in (bar at 7.10)")
                bad += 1
    print(f"{path}: {bad} collision(s)")
    return bad

if __name__ == "__main__":
    for p in sys.argv[1:]: check(p)
