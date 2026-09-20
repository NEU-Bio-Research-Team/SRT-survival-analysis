# -*- coding: utf-8 -*-
"""Approximate renderer + layout linter for the generated decks."""
import sys, io, os
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE

import glob as _g
def _f(pat, fb):
    h = _g.glob(pat, recursive=True)
    return h[0] if h else fb
FD = "/usr/share/fonts/truetype/freefont/"
FONTS = {(0,0): _f("/usr/share/fonts/**/Carlito-Regular.ttf", FD+"FreeSans.ttf"),
         (1,0): _f("/usr/share/fonts/**/Carlito-Bold.ttf", FD+"FreeSansBold.ttf"),
         (0,1): _f("/usr/share/fonts/**/Carlito-Italic.ttf", FD+"FreeSansOblique.ttf"),
         (1,1): _f("/usr/share/fonts/**/Carlito-BoldItalic.ttf", FD+"FreeSansBoldOblique.ttf")}
_cache={}
def font(sz, bold=False, ital=False):
    k=(round(sz),bool(bold),bool(ital))
    if k not in _cache: _cache[k]=ImageFont.truetype(FONTS[(int(bold),int(ital))], max(6,round(sz)))
    return _cache[k]

PPI = 120
SW, SH = int(13.3333*PPI), int(7.5*PPI)
def E(v): return int(Emu(v).inches*PPI) if v is not None else 0

def rgb(c, default=(0,32,96)):
    try:
        if c and c.type is not None and c.rgb is not None:
            return tuple(c.rgb)
    except Exception: pass
    return default

def wrap(draw, text, f, maxw):
    out=[]
    for raw in text.split("\n"):
        words=raw.split(" "); line=""
        for w in words:
            t=(line+" "+w).strip()
            if draw.textlength(t,font=f)<=maxw or not line: line=t
            else: out.append(line); line=w
        out.append(line)
    return out

def draw_tf(draw, tf, x, y, w, h, issues, name, clip=True):
    """Render a text frame; report overflow."""
    cy = y + 2
    maxw = max(10, w-8)
    for p in tf.paragraphs:
        runs=[r for r in p.runs if r.text]
        if not runs:
            cy += 6; continue
        sz = None; bold=False; ital=False; col=(0,32,96)
        for r in runs:
            if r.font.size: sz = r.font.size.pt
            bold = bold or bool(r.font.bold); ital = ital or bool(r.font.italic)
            col = rgb(r.font.color, col)
        sz = sz or 18
        f = font(sz*PPI/72.0, bold, ital)
        txt = "".join(r.text for r in runs)
        lines = wrap(draw, txt, f, maxw)
        lh = sz*PPI/72.0*1.22
        sb = (p.space_before.pt if p.space_before else 0)*PPI/72.0
        sa = (p.space_after.pt if p.space_after else 0)*PPI/72.0
        cy += sb
        for ln in lines:
            tw = draw.textlength(ln, font=f)
            al = str(p.alignment)
            ax = x+4
            if "CENTER" in al: ax = x + (w-tw)/2
            elif "RIGHT" in al: ax = x + w - tw - 4
            draw.text((ax, cy), ln, font=f, fill=col)
            cy += lh
        cy += sa
    if clip and cy > y + h + 6:
        issues.append(f"OVERFLOW  {name}: text needs {(cy-y)/PPI:.2f}in, box is {h/PPI:.2f}in")
    return cy

def render(path, outdir, only=None):
    prs = Presentation(path)
    os.makedirs(outdir, exist_ok=True)
    allissues=[]
    for i, sl in enumerate(prs.slides):
        if only and (i+1) not in only: continue
        img = Image.new("RGB",(SW,SH),(255,255,255))
        d = ImageDraw.Draw(img)
        issues=[]
        for sh in sl.shapes:
            x,y,w,h = E(sh.left),E(sh.top),E(sh.width),E(sh.height)
            nm = sh.name[:26]
            if x < -6 or y < -6 or x+w > SW+6 or y+h > SH+6:
                issues.append(f"OUTSIDE   {nm}: ({x/PPI:.2f},{y/PPI:.2f}) {w/PPI:.2f}x{h/PPI:.2f}in")
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    pim = Image.open(io.BytesIO(sh.image.blob)).convert("RGB").resize((max(1,w),max(1,h)))
                    img.paste(pim,(x,y))
                except Exception as e: d.rectangle([x,y,x+w,y+h], outline=(200,0,0))
                continue
            if sh.has_table:
                tbl = sh.table
                cw=[E(c.width) for c in tbl.columns]; rh=[E(r.height) for r in tbl.rows]
                ty=y
                for ri,row in enumerate(tbl.rows):
                    tx=x
                    for ci,cell in enumerate(row.cells):
                        fill=(255,255,255)
                        try:
                            if cell.fill.type is not None and cell.fill.fore_color.rgb: fill=tuple(cell.fill.fore_color.rgb)
                        except Exception: pass
                        d.rectangle([tx,ty,tx+cw[ci],ty+rh[ri]], fill=fill)
                        endy = draw_tf(d, cell.text_frame, tx+4, ty+3, cw[ci]-8, rh[ri]-6, issues,
                                       f"table r{ri}c{ci} p{i+1}")
                        tx+=cw[ci]
                    ty+=rh[ri]
                if ty > SH-40: issues.append(f"TABLE TOO TALL p{i+1}: ends at {ty/PPI:.2f}in")
                continue
            # autoshape / textbox
            fill=None
            try:
                if sh.fill.type is not None and sh.fill.type == 1: fill=tuple(sh.fill.fore_color.rgb)
            except Exception: pass
            if fill: d.rectangle([x,y,x+w,y+h], fill=fill)
            if sh.has_text_frame and sh.text_frame.text.strip():
                draw_tf(d, sh.text_frame, x, y, w, h, issues, f"{nm} p{i+1}",
                        clip=(fill is None or h>0.5*PPI))
        img.save(f"{outdir}/s{i+1:02d}.png")
        if issues:
            allissues.append((i+1, issues))
    for pg, iss in allissues:
        print(f"\n--- slide {pg} ---")
        for m in iss: print("   ", m)
    if not allissues: print("no layout issues detected")
    return allissues

if __name__ == "__main__":
    render(sys.argv[1], sys.argv[2])
