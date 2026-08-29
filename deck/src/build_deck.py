# -*- coding: utf-8 -*-
import sys, os
sys.path.insert(0, "deck/src")
from deck import *
from content import C
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ACC = [TEAL, ORANGE, PURPLE]
TIER = [MUTED, TEAL, TEAL, NAVY, PURPLE, MUTED]
import json
FT = json.load(open("deck/src/fig_titles.json"))

def reorder_prune(prs, order):
    lst = prs.slides._sldIdLst
    ids = list(lst)
    keep = [ids[i] for i in order]
    for i in ids:
        if i not in keep:
            prs.part.drop_rel(i.rId)
        lst.remove(i)
    for i in keep:
        lst.append(i)

def build(lang):
    T = C[lang]
    F = lambda n: f"deck/fig/{n}_{lang}.png"
    prs = Presentation(TPL)
    AUTHOR = prs.slides[2].shapes[1].text_frame.text  # keep the template's footer name

    # ---------------- 1. title slide -------------------------------------
    s = prs.slides[0]
    tb = s.shapes[1].text_frame; tb.clear()
    p = tb.paragraphs[0]
    r = p.add_run(); r.text = T["deck_title"]
    r.font.name = FONT; r.font.size = Pt(44); r.font.bold = True; r.font.color.rgb = NAVY
    p.alignment = PP_ALIGN.CENTER
    p2 = tb.add_paragraph()
    r2 = p2.add_run(); r2.text = T["deck_sub"]
    r2.font.name = FONT; r2.font.size = Pt(18); r2.font.color.rgb = TEAL
    p2.alignment = PP_ALIGN.CENTER; p2.space_before = Pt(10)
    set_text(s.shapes[4], T["deck_kicker"], size=20, bold=True, color=TEXT, align=PP_ALIGN.CENTER)
    s.shapes[2].top, s.shapes[2].height = Inches(4.12), Inches(1.45)
    gb = s.shapes[2].text_frame; gb.clear()
    for i, line in enumerate(T["group"]):
        pp = gb.paragraphs[0] if i == 0 else gb.add_paragraph()
        rr = pp.add_run(); rr.text = line
        rr.font.name = FONT; rr.font.size = Pt(19 if i < 2 else 15)
        rr.font.bold = i < 2; rr.font.color.rgb = TEXT if i < 2 else MUTED
        pp.alignment = PP_ALIGN.CENTER
    set_text(s.shapes[5], T["tag"], size=13, color=WHITE, align=PP_ALIGN.CENTER)

    # ---------------- 2. outline -----------------------------------------
    s = prs.slides[1]
    tag(s, T["tag"], AUTHOR)
    set_text(s.shapes[4], T["outline_title"], size=32, bold=True, color=TEXT, align=PP_ALIGN.CENTER)
    ob = s.shapes[5]
    ob.left, ob.top, ob.width, ob.height = Inches(2.3), Inches(1.5), Inches(9.4), Inches(4.6)
    tf = ob.text_frame; tf.clear(); tf.word_wrap = True
    for i, it in enumerate(T["outline"]):
        rich(tf, [(f"{i+1}.   ", True, TEAL), (it, False, TEXT)],
             size=21, first=(i == 0), space_after=13)

    N = 0
    def new(title=None):
        nonlocal N
        sl = clone(prs, 2); N += 1
        tag(sl, T["tag"], AUTHOR)
        if title is not None: set_title(sl, title)
        return sl

    def sec(i):
        nonlocal N
        num, ttl, sub = T["sections"][i]
        section(prs, num, ttl, sub, T["tag"], AUTHOR); N += 1

    def figslide(title, fig, notetxt, top=1.02, bottom=6.42):
        if title is None:
            title = FT[lang][fig.split("_")[0] + "_title"]
        sl = new(title); picture(sl, F(fig), top=top, bottom=bottom)
        if notetxt: note(sl, notetxt, top=6.50)
        return sl

    # ================= SECTION 1 ==========================================
    sec(0)
    figslide(None, "f1_motivation", T["s4_note"], bottom=6.42)

    s = new(T["s5_title"])
    table(s, T["s5_tbl"], 0.85, 1.12, 11.6, col_w=[1, 1.25], size=15, head_size=15,
          row_h=0.62, head_h=0.44)
    tb = box(s, 1.2, 4.55, 10.9, 0.9)
    tf = tb.text_frame; tf.word_wrap = True
    para(tf, "“" + T["s5_quote"] + "”", size=22, bold=True, color=NAVY,
         align=PP_ALIGN.CENTER, first=True, italic=True)
    rule(s, 4.4, 5.62, 4.5, RULE)
    tb2 = box(s, 0.7, 5.80, 11.9, 0.6)
    tf2 = tb2.text_frame; tf2.word_wrap = True
    para(tf2, T["s5_chain"], size=13.5, color=TEAL, align=PP_ALIGN.CENTER, first=True)

    # ================= SECTION 2 ==========================================
    sec(1)
    figslide(None, "f2_timeline", T["s7_note"], bottom=6.38)

    s = new(T["s8_title"])
    table(s, T["s8_tbl"], 0.62, 1.10, 12.1, col_w=[2.5, 3.4, 2.1, 1.6], size=14,
          head_size=13.5, row_h=0.46, head_h=0.40)
    card(s, 0.62, 3.35, 5.95, 1.95, T["s8_c1h"], T["s8_c1"], TEAL, 15, 13.5)
    card(s, 6.77, 3.35, 5.95, 1.95, T["s8_c2h"], T["s8_c2"], ORANGE, 15, 13.5)

    s = new(T["s9_title"])
    kpis(s, T["s9_kpi"], top=1.12, h=1.25, size_v=32, size_l=12.5)
    for i, (h, b, a) in enumerate(T["s9_cards"]):
        card(s, 0.62 + i * 4.09, 2.85, 3.93, 2.15, h, b, ACC[a], 15, 13)

    # ================= SECTION 3 ==========================================
    sec(2)
    s = new(T["s11_title"])
    y = 1.06
    for ti, (tn, name, sub, files, a) in enumerate(T["s11_tiers"]):
        bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.62), Inches(y), Inches(12.1), Inches(0.80))
        bg.fill.solid(); bg.fill.fore_color.rgb = BAND
        bg.line.color.rgb = RULE; bg.line.width = Pt(0.75); bg.shadow.inherit = False
        chip(s, 0.62, y, 1.30, 0.80, tn, NAVY, WHITE, 13, True, radius=False)
        tb = box(s, 2.06, y + 0.07, 2.75, 0.70)
        tf = tb.text_frame; tf.word_wrap = True
        para(tf, name, size=14, bold=True, color=TEXT, first=True, space_after=0)
        para(tf, sub, size=10.5, color=MUTED, space_after=0)
        x = 4.95
        col = TIER[ti]
        for j, fn in enumerate(files):
            w = min(0.13 + len(fn) * 0.083, 2.55)
            if x + w > 12.40:
                chip(s, x, y + 0.23, 0.42, 0.34, "…", col, WHITE, 12, radius=True)
                break
            chip(s, x, y + 0.23, w, 0.34, fn, col, WHITE, 10.5, radius=True)
            x += w + 0.10
        y += 0.86
    note(s, T["s11_note"], top=6.34, size=13)

    s = new(T["s12_title"])
    table(s, T["s12_tbl"], 0.62, 1.10, 12.1, col_w=[0.5, 3.2, 5.3, 3.5], size=12.5,
          head_size=12.5, row_h=0.56, head_h=0.40, aligns=["c", "l", "l", "l"])
    card(s, 0.62, 5.02, 12.1, 1.20, T["s12_note_h"], [T["s12_note"]], ORANGE, 14, 13.5)

    s = new(T["s13_title"])
    table(s, T["s13_tbl"], 0.62, 1.02, 12.1, col_w=[2.3, 2.9, 0.75, 2.1, 4.0], size=11,
          head_size=11, row_h=0.40, head_h=0.36, aligns=["l", "l", "c", "l", "l"])

    figslide(T["s14_title"], "f6_blocks", T["s14_note"], top=1.45, bottom=5.90)

    # ================= SECTION 4 ==========================================
    sec(3)
    figslide(None, "f3_km", T["s16_note"], bottom=6.42)
    figslide(None, "f4_duration", T["s17_note"], bottom=6.42)
    figslide(None, "f5_threshold", T["s18_note"], bottom=6.42)
    figslide(None, "f11_eventrate", T["s19_note"], bottom=6.38)
    figslide(None, "f9_importers", T["s20_note"], bottom=6.42)
    figslide(None, "f12_valdur", T["s21_note"], bottom=6.42)
    figslide(None, "f7_scope", T["s22_note"], top=1.40, bottom=6.05)
    figslide(None, "f10_hazard", T["s23_note"], bottom=6.38)

    # ================= SECTION 5 ==========================================
    sec(4)
    figslide(None, "f8_quality", T["s25_note"], top=1.15, bottom=6.15)

    s = new(T["s26_title"])
    pos = [(0.62, 1.08), (0.62, 2.82), (0.62, 4.56), (6.77, 1.08), (6.77, 2.82)]
    for (h, b, a), (x, yy) in zip(T["s26"], pos):
        card(s, x, yy, 5.95, 1.55, h, b, ACC[a], 14, 12)
    card(s, 6.77, 4.56, 5.95, 1.55, T["s26_rule_h"], [T["s26_rule"]], NAVY, 14, 12)

    # ================= SECTION 6 ==========================================
    sec(5)
    s = new(T["s28_title"])
    table(s, T["s28_tbl"], 0.62, 1.12, 12.1, col_w=[1.7, 3.3, 3.6, 3.5], size=13,
          head_size=13, row_h=0.68, head_h=0.42)
    note(s, T["s28_note"], top=5.05, size=14, color=NAVY, italic=True)

    s = new(T["s29_title"])
    for i, (h, b) in enumerate(T["s29"]):
        x = 0.62 + (i % 3) * 4.09
        yy = 1.28 + (i // 3) * 2.30
        card(s, x, yy, 3.93, 2.05, h, b, ACC[i % 3], 14, 12.5)

    # ---------------- thank-you ------------------------------------------
    ty = prs.slides[3]
    tag(ty, T["tag"], AUTHOR)
    ty.shapes[4].left, ty.shapes[4].width = Inches(0), Inches(SW)
    set_text(ty.shapes[4], T["thanks"], size=44, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    tb = box(ty, 1.6, 3.60, 10.1, 0.8)
    tf = tb.text_frame; tf.word_wrap = True
    para(tf, T["thanks_sub"], size=14, color=MUTED, align=PP_ALIGN.CENTER, first=True)

    order = [0, 1] + list(range(4, 4 + N)) + [3]
    reorder_prune(prs, order)
    out = f"deck/out/Sinking_Relationships_Data_{'VI' if lang=='vi' else 'EN'}.pptx"
    prs.save(out)
    print(f"saved {out}  ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)")
    return out

if __name__ == "__main__":
    for lg in ("vi", "en"):
        build(lg)
