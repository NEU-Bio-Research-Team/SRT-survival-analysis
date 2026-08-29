# -*- coding: utf-8 -*-
import sys; sys.path.insert(0, "deck/src")
from deck2 import *
from deck2 import C as _hex
from content2 import C
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

TIER_COL = [MUTED, BLUE, BLUE_M, BLUE_D]

def reorder_prune(prs, order):
    lst = prs.slides._sldIdLst; ids = list(lst)
    keep = [ids[i] for i in order]
    for i in ids:
        if i not in keep: prs.part.drop_rel(i.rId)
        lst.remove(i)
    for i in keep: lst.append(i)

def build(lang):
    T = C[lang]
    F = lambda n: f"deck/fig2/{n}_{lang}.png"
    prs = Presentation(TPL)
    AU = prs.slides[2].shapes[1].text_frame.text
    N = 0

    # ---------------- title -------------------------------------------
    s = prs.slides[0]
    tf = s.shapes[1].text_frame; tf.clear()
    p = tf.paragraphs[0]; r = p.add_run(); r.text = T["title"]
    r.font.name = FONT; r.font.size = Pt(46); r.font.bold = True; r.font.color.rgb = NAVY
    p.alignment = PP_ALIGN.CENTER
    p2 = tf.add_paragraph(); r2 = p2.add_run(); r2.text = T["sub"]
    r2.font.name = FONT; r2.font.size = Pt(17); r2.font.color.rgb = BLUE
    p2.alignment = PP_ALIGN.CENTER; p2.space_before = Pt(12)
    tk = s.shapes[4].text_frame; tk.clear()
    para(tk, T["kicker"], size=18, bold=True, color=BODY, align=PP_ALIGN.CENTER,
         first=True, after=0)
    s.shapes[2].top, s.shapes[2].height = Inches(4.06), Inches(1.24)
    gb = s.shapes[2].text_frame; gb.clear()
    for i, line in enumerate(T["group"]):
        para(gb, line, size=17 if i < 2 else 14, bold=(i < 2),
             color=INK if i < 2 else MUTED, align=PP_ALIGN.CENTER, first=(i == 0), after=2)
    nb = box(s, 2.24, 5.34, 8.67, 0.36)
    para(nb.text_frame, T["group_note"], size=13, color=MUTED,
         align=PP_ALIGN.CENTER, first=True, after=0)
    tg = s.shapes[5].text_frame; tg.clear()
    para(tg, T["tag"], size=12, color=WHITE, align=PP_ALIGN.CENTER, first=True, after=0)

    # ---------------- outline -----------------------------------------
    s = prs.slides[1]
    chrome(s, T["tag"], AU, T["outline_t"])
    body = [x for x in s.shapes if x.has_text_frame and "About/Syllabus" in x.text_frame.text
            or (x.has_text_frame and x.text_frame.text.strip().startswith(("About", "Group Req")))]
    for x in body: x._element.getparent().remove(x._element)
    for i, (h, d) in enumerate(T["outline"]):
        col = 0 if i < 3 else 1
        row = i % 3
        x = LEFT + col * 6.30; y = 1.55 + row * 1.62
        rule(s, x, y, 5.60, BLUE_M, 0.028)
        tb = box(s, x, y + 0.14, 5.60, 1.20)
        tf = tb.text_frame
        rich(tf, [(f"{i+1}    ", True, BLUE_M), (h, True, INK)], size=19, first=True, after=4)
        para(tf, d, size=13.5, color=MUTED, after=0)

    def sec(i):
        nonlocal N
        n, t, sb = T["sections"][i]
        section_slide(prs, T["tag"], AU, n, t, sb); N += 1

    def new(title, lead=None):
        nonlocal N
        N += 1
        return content_slide(prs, T["tag"], AU, title, lead)

    def figs(title, lead, fig, n=None, top=1.52, bottom=6.86, mw=12.0):
        s = new(title, lead)
        picture(s, F(fig), top, min(bottom, 6.44) if n else min(bottom, 6.92), mw)
        if n: note(s, n)
        return s

    # ================= 01 ==============================================
    sec(0)
    figs(T["s_story_t"], T["s_story_l"], "story", T["s_story_n"], bottom=6.80)

    s = new(T["s_why_t"])
    rule_table(s, T["s_why_tbl"], LEFT, 1.30, CW, col_w=[1, 1.3], size=14.5,
               head_size=14.5, row_h=0.66, head_h=0.46)
    qb = box(s, LEFT, 4.68, CW, 0.9)
    para(qb.text_frame, T["s_why_q"], size=23, bold=True, color=BLUE_D,
         align=PP_ALIGN.CENTER, first=True, after=0)
    callout(s, LEFT, 5.72, CW, T["s_why_c"][0], T["s_why_c"][1], BLUE_D, TINT_B,
            bh=0.34, body_h=0.56)

    # ================= 02 ==============================================
    sec(1)
    s = figs(T["s_unit_t"], T["s_unit_l"], "timeline", None, top=1.46, bottom=4.94, mw=11.6)
    callout(s, LEFT, 5.16, 6.05, T["s_unit_c1"][0], T["s_unit_c1"][1], BLUE_D, TINT_B,
            bh=0.34, body_h=0.86)
    callout(s, 7.29, 5.16, 6.05, T["s_unit_c2"][0], T["s_unit_c2"][1], ORANGE, TINT_O,
            bh=0.34, body_h=0.86)

    s = new(T["s_files_t"], T["s_files_l"])
    rule_table(s, T["s_files_tbl"], LEFT, 1.66, CW, col_w=[2.0, 3.9, 3.6, 2.0],
               size=13.5, head_size=13.5, row_h=0.62, head_h=0.44)
    callout(s, LEFT, 4.66, CW, T["s_files_c"][0], T["s_files_c"][1], BLUE_D, TINT_B,
            bh=0.34, body_h=0.80)

    s = new(T["s_size_t"], T["s_size_l"])
    kpi_row(s, T["s_size_kpi"], t=1.62, h=1.10)
    h1, b1 = T["s_size_b1"]; h2, b2 = T["s_size_b2"]
    bullets(s, LEFT, 3.20, 6.05, 3.0, h1, b1, size=14, head_size=15.5)
    bullets(s, 7.29, 3.20, 6.05, 3.0, h2, b2, size=14, head_size=15.5)

    # ================= 03 ==============================================
    sec(2)
    s = new(T["s_map_t"], T["s_map_l"])
    y = 1.62
    for tn, name, sub, ci in T["s_map_tiers"]:
        col = TIER_COL[ci]
        rule(s, LEFT, y, CW, RULE, 0.008)
        tb = box(s, LEFT, y + 0.10, 1.55, 0.4)
        para(tb.text_frame, tn, size=13, bold=True, color=col, first=True, after=0)
        nb2 = box(s, 2.20, y + 0.06, 3.85, 0.5)
        para(nb2.text_frame, name, size=15, bold=True, color=INK, first=True, after=0)
        db = box(s, 6.15, y + 0.09, 6.55, 0.5)
        para(db.text_frame, sub, size=13.5, color=BODY, first=True, after=0)
        y += 0.80
    rule(s, LEFT, y, CW, RULE, 0.008)
    note(s, T["s_map_n"], t=y + 0.18, size=13)

    s = new(T["s_lvl_t"], T["s_lvl_l"])
    rule_table(s, T["s_lvl_tbl"], LEFT, 1.66, CW, col_w=[3.1, 5.0, 4.3],
               size=13, head_size=13, row_h=0.72, head_h=0.42)
    callout(s, LEFT, 5.86, CW, T["s_lvl_c"][0], T["s_lvl_c"][1], ORANGE, TINT_O,
            bh=0.34, body_h=0.58)

    figs(T["s_blk_t"], T["s_blk_l"], "blocks", T["s_blk_n"], top=1.70, bottom=6.60, mw=12.2)

    # ================= 04 ==============================================
    sec(3)
    figs(T["s_flows_t"], T["s_flows_l"], "flows", T["s_flows_n"], top=1.50, bottom=6.86)
    figs(T["s_km_t"], T["s_km_l"], "km", T["s_km_n"], bottom=6.80, mw=11.4)
    figs(T["s_dur_t"], T["s_dur_l"], "duration", T["s_dur_n"], bottom=6.78, mw=11.4)
    figs(T["s_sz_t"], T["s_sz_l"], "size", T["s_sz_n"], bottom=6.80, mw=11.4)
    figs(T["s_sec_t"], T["s_sec_l"], "sectors", T["s_sec_n"], bottom=6.82)
    figs(T["s_map2_t"], T["s_map2_l"], "map", T["s_map2_n"], bottom=6.82, mw=11.5)
    figs(T["s_mk_t"], T["s_mk_l"], "markets", T["s_mk_n"], bottom=6.80, mw=11.4)
    figs(T["s_cc_t"], T["s_cc_l"], "conc", T["s_cc_n"], bottom=6.80, mw=11.6)
    figs(T["s_hz_t"], T["s_hz_l"], "hazard", T["s_hz_n"], bottom=6.80, mw=11.0)

    # ================= 05 ==============================================
    sec(4)
    figs(T["s_q_t"], T["s_q_l"], "quality", T["s_q_n"], top=1.60, bottom=6.76)

    s = new(T["s_lim_t"])
    pos = [(LEFT, 1.34), (LEFT, 3.04), (LEFT, 4.74), (7.29, 1.34), (7.29, 3.04)]
    cols = [(BLUE_D, TINT_B), (ORANGE, TINT_O), (BLUE, TINT_B)]
    for (h, b, ci), (x, yy) in zip(T["s_lim"], pos):
        c1, c2 = cols[ci]
        callout(s, x, yy, 6.05, h, b, c1, c2, bh=0.36, body_h=1.18, size=13)
    callout(s, 7.29, 4.74, 6.05, T["s_lim_c"][0], T["s_lim_c"][1], MUTED, _hex("E8EEF3"),
            bh=0.36, body_h=1.18, size=13)

    # ================= 06 ==============================================
    sec(5)
    figs(T["s_sc_t"], T["s_sc_l"], "scope", T["s_sc_n"], top=1.86, bottom=6.10)

    s = new(T["s_st_t"])
    rule_table(s, T["s_st_tbl"], LEFT, 1.40, CW, col_w=[3.1, 5.4, 3.9],
               size=14.5, head_size=14.5, row_h=0.74, head_h=0.46)
    callout(s, LEFT, 5.10, CW, T["s_st_c"][0], T["s_st_c"][1], BLUE_D, TINT_B,
            bh=0.34, body_h=0.56)

    s = new(T["s_dec_t"])
    for i, (h, b) in enumerate(T["s_dec"]):
        x = LEFT + (i % 3) * 4.16
        yy = 1.46 + (i // 3) * 2.62
        rule(s, x, yy, 3.86, BLUE_M, 0.028)
        tb = box(s, x, yy + 0.16, 3.86, 2.20)
        tf = tb.text_frame
        rich(tf, [(f"{i+1:02d}    ", True, BLUE_M), (h, True, INK)], size=15.5,
             first=True, after=7)
        para(tf, b, size=13.5, color=BODY, after=0)

    # ---------------- thanks ------------------------------------------
    ty = prs.slides[3]
    ref = list(ty.shapes)
    backdrop(ty)
    sh = ref[3]; sh.width = Inches(3.6)
    tfx = sh.text_frame; tfx.clear(); para(tfx, T["tag"], size=13, color=WHITE, first=True, after=0)
    tfa = ref[1].text_frame; tfa.clear(); para(tfa, AU, size=13, color=WHITE, first=True, after=0)
    t4 = ref[4]
    t4.left, t4.top, t4.width, t4.height = Inches(0), Inches(2.85), Inches(SW), Inches(0.9)
    tt = t4.text_frame; tt.clear()
    para(tt, T["thanks"], size=42, bold=True, color=INK, align=PP_ALIGN.CENTER,
         first=True, after=0)
    rule(ty, 5.17, 4.02, 3.0, BLUE_M, 0.030)
    sb = box(ty, 1.6, 4.24, 10.1, 0.6)
    para(sb.text_frame, T["thanks_sub"], size=14, color=MUTED,
         align=PP_ALIGN.CENTER, first=True, after=0)

    reorder_prune(prs, [0, 1] + list(range(4, 4 + N)) + [3])
    out = f"deck/out/Sinking_Relationships_{'VI' if lang == 'vi' else 'EN'}.pptx"
    prs.save(out); print(f"saved {out}")
    return out

if __name__ == "__main__":
    for lg in ("vi", "en"): build(lg)
